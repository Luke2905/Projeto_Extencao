from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


PROJECT_ROOT = Path(__file__).resolve().parent
OUTPUT_FILE = PROJECT_ROOT / "Documentacao_ProjetoExtensao.docx"


def set_document_defaults(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    style.font.size = Pt(11)


def add_title_page(document: Document) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Documentação Técnica do Projeto de Extensão")
    run.bold = True
    run.font.size = Pt(20)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Sistema de Monitoramento de Represas (Java + MySQL)")
    run.font.size = Pt(13)

    date_paragraph = document.add_paragraph()
    date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_paragraph.add_run(
        f"Versão do documento: 1.0\nGerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    )

    document.add_page_break()


def add_intro(document: Document) -> None:
    document.add_heading("1. Visão Geral", level=1)
    document.add_paragraph(
        "Este projeto foi desenvolvido no contexto de Extensão Curricular para integrar "
        "conhecimentos de Banco de Dados, Programação Java, Estatística e análise ambiental. "
        "A aplicação desktop permite consultar dados de monitoramento de represas armazenados "
        "em um banco MySQL e apresentar resultados estatísticos por meio de uma interface Swing."
    )

    document.add_heading("2. Objetivos", level=1)
    for item in [
        "Centralizar a consulta de dados históricos de monitoramento de represas.",
        "Disponibilizar indicadores rápidos para apoio à análise ambiental.",
        "Facilitar a visualização dos registros por meio de tabela em interface gráfica.",
        "Apoiar a tomada de decisão com métricas como total de registros, média histórica e menor volume observado.",
    ]:
        document.add_paragraph(item, style="List Bullet")


def add_architecture(document: Document) -> None:
    document.add_heading("3. Arquitetura da Solução", level=1)
    document.add_paragraph(
        "A aplicação segue uma estrutura simples em camadas leves: entrada da aplicação "
        "(classe Principal), acesso a dados (Conexao e DataBase) e apresentação gráfica "
        "(JOptionPane/JTable e classe Tela como alternativa de janela dedicada)."
    )

    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "Camada"
    header[1].text = "Classe(s)"
    header[2].text = "Responsabilidade"

    rows = [
        (
            "Entrada",
            "Principal",
            "Iniciar o programa, exibir menu e delegar ações para DataBase.",
        ),
        (
            "Conexão",
            "Conexao",
            "Gerenciar parâmetros JDBC e abrir conexão com MySQL.",
        ),
        (
            "Dados/Serviço",
            "DataBase",
            "Executar consultas SQL, transformar resultados e exibir informações.",
        ),
        (
            "Interface",
            "JOptionPane, JTable, Tela",
            "Exibir tabela e mensagens para interação com o usuário.",
        ),
    ]

    for camada, classes, responsabilidade in rows:
        row = table.add_row().cells
        row[0].text = camada
        row[1].text = classes
        row[2].text = responsabilidade


def add_structure(document: Document) -> None:
    document.add_heading("4. Estrutura de Pastas", level=1)
    document.add_paragraph(
        "Estrutura identificada no repositório local (Projeto Eclipse Java):"
    )
    code = document.add_paragraph()
    code.style = "No Spacing"
    code.add_run(
        ".\n"
        "├─ src/projetoextensao/\n"
        "│  ├─ Principal.java\n"
        "│  ├─ Conexao.java\n"
        "│  ├─ DataBase.java\n"
        "│  └─ Tela.java\n"
        "├─ bin/projetoextensao/*.class\n"
        "├─ .classpath\n"
        "├─ .project\n"
        "└─ README.md\n"
    )


def add_class_details(document: Document) -> None:
    document.add_heading("5. Descrição Detalhada das Classes", level=1)

    document.add_heading("5.1 Principal.java", level=2)
    document.add_paragraph(
        "É o ponto de entrada da aplicação. Exibe uma mensagem inicial e apresenta "
        "um menu interativo com as opções de consulta."
    )
    for item in [
        "Opção 1: lista todos os dados em tabela.",
        "Opção 2: exibe o total de registros existentes.",
        "Opção 3: exibe a média histórica de chuva (mm).",
        "Opção 4: exibe o menor volume útil registrado.",
        "Opção 0: encerra o programa.",
    ]:
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("5.2 Conexao.java", level=2)
    document.add_paragraph(
        "Centraliza os parâmetros de conexão JDBC com MySQL e fornece o método "
        "`getConect()` para abrir conexão. Em caso de erro, imprime a mensagem no stderr "
        "e retorna `null`."
    )
    document.add_paragraph(
        "Parâmetros atuais no código:\n"
        "- URL: jdbc:mysql://localhost:3306/monitoramento_represa\n"
        "- Usuário: root\n"
        "- Senha: vazia"
    )

    document.add_heading("5.3 DataBase.java", level=2)
    document.add_paragraph(
        "Concentra as consultas SQL e a montagem da interface de saída."
    )
    funcs = [
        ("listarDados()", "Executa JOIN entre MONITORAMENTO e REPRESA e mostra JTable."),
        ("totalDados()", "Executa COUNT(*) e mostra total de registros."),
        ("mediaDados()", "Executa AVG(chuva_mm) e mostra média histórica."),
        ("menorRegistro()", "Busca o menor volume útil (ORDER BY ASC LIMIT 1)."),
    ]
    for nome, desc in funcs:
        document.add_paragraph(f"{nome}: {desc}", style="List Bullet")

    document.add_heading("5.4 Tela.java", level=2)
    document.add_paragraph(
        "Define uma janela Swing (`JFrame`) para exibir a tabela e um painel de dashboard. "
        "No estado atual do projeto, essa tela está implementada, mas não está ativa no fluxo "
        "principal (há chamadas comentadas em `DataBase.listarDados()`)."
    )


def add_database_section(document: Document) -> None:
    document.add_heading("6. Banco de Dados", level=1)
    document.add_paragraph(
        "O sistema depende do banco MySQL `monitoramento_represa` e das tabelas "
        "`MONITORAMENTO` e `REPRESA`."
    )

    document.add_heading("6.1 Tabelas esperadas", level=2)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    head = table.rows[0].cells
    head[0].text = "Tabela"
    head[1].text = "Colunas utilizadas no código"

    table_data = [
        (
            "MONITORAMENTO",
            "id, id_represa, data, chuva_mm, volume_util_percent, chuva_acumulada_mes_mm",
        ),
        ("REPRESA", "id, nome"),
    ]
    for tabela, colunas in table_data:
        row = table.add_row().cells
        row[0].text = tabela
        row[1].text = colunas

    document.add_heading("6.2 Consultas SQL implementadas", level=2)
    for sql_item in [
        "Listagem completa com JOIN e ordenação por data decrescente.",
        "Contagem total de registros (COUNT).",
        "Média histórica de chuva em milímetros (AVG).",
        "Menor volume útil registrado, com nome da represa e data.",
    ]:
        document.add_paragraph(sql_item, style="List Bullet")


def add_setup_and_run(document: Document) -> None:
    document.add_heading("7. Requisitos e Configuração", level=1)
    reqs = [
        "Java JDK 24 (conforme configuração do projeto Eclipse).",
        "MySQL Server em execução.",
        "Banco `monitoramento_represa` criado e populado.",
        "Driver JDBC: mysql-connector-j (ex.: 9.6.0) disponível no classpath.",
        "IDE Eclipse (opcional, mas recomendada para este projeto).",
    ]
    for req in reqs:
        document.add_paragraph(req, style="List Bullet")

    document.add_heading("8. Como Executar", level=1)
    document.add_heading("8.1 Via Eclipse", level=2)
    for step in [
        "Importar o projeto Java existente.",
        "Confirmar JRE/JDK do projeto para JavaSE-24.",
        "Adicionar o JAR do MySQL Connector no Build Path.",
        "Executar a classe `projetoextensao.Principal`.",
    ]:
        document.add_paragraph(step, style="List Number")

    document.add_heading("8.2 Via terminal (PowerShell)", level=2)
    document.add_paragraph("Comandos de exemplo:")
    compile_block = document.add_paragraph()
    compile_block.style = "No Spacing"
    compile_block.add_run(
        'javac -encoding UTF-8 -cp ".;lib\\mysql-connector-j-9.6.0.jar" '
        "-d bin src\\projetoextensao\\*.java"
    )
    run_block = document.add_paragraph()
    run_block.style = "No Spacing"
    run_block.add_run(
        'java -cp "bin;lib\\mysql-connector-j-9.6.0.jar" projetoextensao.Principal'
    )


def add_risks_improvements(document: Document) -> None:
    document.add_heading("9. Limitações Atuais e Melhorias Recomendadas", level=1)
    for point in [
        "Credenciais de banco estão fixas no código; ideal usar variáveis de ambiente ou arquivo de configuração.",
        "Métodos de consulta não validam retorno nulo da conexão antes de usar `prepareStatement`.",
        "A média é recuperada como inteiro (`getInt`), podendo perder precisão decimal; preferir `getDouble`.",
        "Há import não utilizado em `Principal` (`SwitchCase`) e classe `Tela` ainda não integrada ao fluxo principal.",
        "Falta camada de testes automatizados (unitários/integrados).",
        "Falta gerenciamento de dependências (Maven/Gradle) para facilitar reprodução do ambiente.",
    ]:
        document.add_paragraph(point, style="List Bullet")

    document.add_heading("10. Fluxo de Uso do Sistema", level=1)
    for step in [
        "Usuário inicia a aplicação (`Principal`).",
        "Sistema apresenta menu de opções no `JOptionPane`.",
        "Conforme opção escolhida, `DataBase` executa a consulta SQL apropriada.",
        "Resultados são exibidos ao usuário em tabela ou mensagem.",
        "Usuário pode repetir consultas até selecionar saída.",
    ]:
        document.add_paragraph(step, style="List Number")

    document.add_heading("11. Manutenção", level=1)
    document.add_paragraph(
        "Para manutenção evolutiva, recomenda-se priorizar: externalização de configuração, "
        "tratamento robusto de erros de conexão, refatoração para camadas (DAO/Service/UI), "
        "adoção de build com Maven/Gradle e criação de testes automatizados."
    )


def add_appendix(document: Document) -> None:
    document.add_heading("12. Apêndice: Dados Técnicos do Projeto", level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    head = table.rows[0].cells
    head[0].text = "Item"
    head[1].text = "Valor identificado"

    values = [
        ("Nome do projeto Eclipse", "ProjetoExtensao"),
        ("Pacote Java", "projetoextensao"),
        ("Versão de compilação", "Java 24"),
        ("Pasta de fontes", "src"),
        ("Pasta de saída", "bin"),
        ("Driver JDBC configurado", "mysql-connector-j-9.6.0.jar"),
    ]
    for item, value in values:
        row = table.add_row().cells
        row[0].text = item
        row[1].text = value


def main() -> None:
    document = Document()
    set_document_defaults(document)
    add_title_page(document)
    add_intro(document)
    add_architecture(document)
    add_structure(document)
    add_class_details(document)
    add_database_section(document)
    add_setup_and_run(document)
    add_risks_improvements(document)
    add_appendix(document)
    document.save(OUTPUT_FILE)
    print(f"Arquivo gerado com sucesso: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
